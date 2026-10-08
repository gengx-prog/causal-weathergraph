"""Build the complete bilingual response, editable DOCX, and coverage receipt.

PDF export is a separate Word operation so the rendered artifact can be inspected.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PARTS = ROOT / 'revision/response_parts'
EXPECTED = (['ED-' + str(i) for i in range(1, 7)]
            + ['R1-' + p + str(i) for p, count in [('A', 5), ('I', 6), ('M', 7),
               ('P', 10), ('D', 6), ('REF', 1), ('O', 1), ('T', 5), ('L', 3)]
               for i in range(1, count + 1)]
            + ['R2-' + str(i) for i in range(1, 6)]
            + ['R3-' + str(i) for i in range(1, 8)])
def manuscript_highlights():
    """Keep the response receipt tied to the actual current manuscript."""
    from pylatexenc.latex2text import LatexNodes2Text
    tex = (ROOT / 'manuscript/revised/Causal_WeatherGraph_elsarticle_DRAFT.tex').read_text(encoding='utf8')
    block = re.search(r'\\begin\{highlights\}(.*?)\\end\{highlights\}', tex, re.S).group(1)
    return [LatexNodes2Text().latex_to_text(t).strip() for t in re.split(r'\\item\s*', block)[1:]]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def set_font(style, size, english='Times New Roman', chinese='宋体', bold=False, color=None):
    style.font.name = english
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), chinese)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = ' '
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for node in [begin, code, separate, text, end]: run._r.append(node)


def create_reference(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(.9); sec.footer_distance = Cm(.9)
    normal = doc.styles['Normal']
    set_font(normal, 10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 24), ('Heading 1', 19), ('Heading 2', 16), ('Heading 3', 12.5), ('Heading 4', 11)]:
        style = doc.styles[name]
        set_font(style, size, 'Arial', '微软雅黑', True, '183D58')
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    for name in ['Quote', 'Intense Quote']:
        set_font(doc.styles[name], 10)
        doc.styles[name].paragraph_format.left_indent = Cm(.45)
        doc.styles[name].paragraph_format.right_indent = Cm(.3)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.core_properties.title = '审稿意见逐条答复与技术证据说明'
    doc.core_properties.subject = 'INS-D-26-9116 Major Revision'
    doc.core_properties.author = ''
    doc.core_properties.comments = 'Evidence-based response document; pending manuscript actions are disclosed.'
    doc.save(path)


def style_document(path):
    doc = Document(path)
    # The source title is the cover title, not a numbered scientific section.
    for p in doc.paragraphs:
        if p.text.strip() == '审稿意见逐条答复与技术证据说明':
            p.style = doc.styles['Title']
            break
    chapter = next(p for p in doc.paragraphs if p.style.name == 'Heading 2')
    title = chapter.insert_paragraph_before('目录')
    title.style = doc.styles['Title']
    title.paragraph_format.page_break_before = True
    toc = chapter.insert_paragraph_before()
    add_field(toc, ' TOC \\o "2-3" \\h \\z \\u ')
    for p in doc.paragraphs:
        if p.style.name == 'Heading 2':
            p.paragraph_format.page_break_before = True
        if p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next = True
        if p.text.strip() in {'审稿原文','答复（英文）','技术实现与证据','证明范围与限制','稿件落实状态'}:
            p.paragraph_format.keep_with_next = True
        if p.style.name == 'Source Code':
            for run in p.runs:
                run.font.size = Pt(8.5)
    for table in doc.tables:
        table.autofit = False
        count = len(table.columns)
        if count == 3:
            widths = [3.0, 7.0, 6.8]
        elif count == 2:
            widths = [3.6, 13.2]
        else:
            widths = [16.8 / count] * count
        for row_num, row in enumerate(table.rows):
            pr = row._tr.get_or_add_trPr()
            block = OxmlElement('w:cantSplit'); pr.append(block)
            if row_num == 0:
                repeat = OxmlElement('w:tblHeader'); pr.append(repeat)
            for idx, cell in enumerate(row.cells):
                cell.width = Cm(widths[min(idx, len(widths)-1)])
                if row_num == 0:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EFF4')
                    cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = 1.05
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if row_num == 0: run.bold = True
    for sec in doc.sections:
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.text = 'INS-D-26-9116  |  审稿意见逐条答复与技术证据'
        for run in hp.runs:
            run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string('657480')
        fp = sec.footer.paragraphs[0]; fp.alignment = 1
        fp.add_run('第 '); add_field(fp, ' PAGE '); fp.add_run(' 页 / '); add_field(fp, ' NUMPAGES '); fp.add_run(' 页')
        for run in fp.runs: run.font.size = Pt(8)
    settings = doc.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None: update = OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'), 'true')
    doc.save(path)


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--output', type=Path, default=ROOT.parent / '审稿答复文档')
    args = p.parse_args(); args.output.mkdir(parents=True, exist_ok=True)
    highlights = manuscript_highlights()
    paths = [PARTS / name for name in ['front.md', 'editor.md', 'reviewer1.md', 'reviewer2.md', 'reviewer3.md', 'appendices.md']]
    texts = [f.read_text(encoding='utf8') for f in paths]
    complete = '\n\n'.join(t.strip() for t in texts) + '\n'
    # A per-item structural audit complements the independent content review.
    matches = list(re.finditer(r'^###\s+((?:ED-\d+)|(?:R1-[A-Z]+\d+)|(?:R[23]-\d+))\s+(.+)$', complete, re.M))
    ids = [m.group(1) for m in matches]
    assert ids == EXPECTED, dict(expected=EXPECTED, actual=ids)
    required = ['审稿原文', '答复', '技术实现与证据', '证明范围与限制', '稿件落实状态']
    coverage = []
    for n, match in enumerate(matches):
        end = matches[n+1].start() if n+1 < len(matches) else len(complete)
        block = complete[match.end():end]
        missing = [label for label in required if label not in block]
        assert not missing, (match.group(1), missing)
        coverage.append(dict(response_id=match.group(1), title=match.group(2), source_quote=True,
                             english_response=True, technical_evidence=True, limits=True,
                             manuscript_status=True, source_characters=len(block)))
    for text in highlights:
        assert len(text) <= 80, (text, len(text))
    stem = '审稿意见逐条答复与技术证据说明'
    markdown = args.output / (stem + '.md')
    markdown.write_text(complete, encoding='utf8')
    (ROOT / 'revision/response_full.md').write_text(complete, encoding='utf8')
    with (args.output/'逐条覆盖核查.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer=csv.DictWriter(handle, fieldnames=coverage[0].keys()); writer.writeheader(); writer.writerows(coverage)
    reference = args.output / 'document_style.docx'
    create_reference(reference)
    docx = args.output / (stem + '.docx')
    subprocess.run(['pandoc', str(markdown), '--from=markdown+tex_math_dollars+tex_math_single_backslash', '--to=docx',
                    '--reference-doc=' + str(reference), '--output=' + str(docx)], check=True, cwd=ROOT)
    style_document(docx)
    manifest = dict(built_utc=datetime.now(timezone.utc).isoformat(), response_count=len(coverage),
                    counts=dict(editor=6, reviewer1=44, reviewer2=5, reviewer3=7),
                    sources={str(f.relative_to(ROOT)):digest(f) for f in paths},
                    highlights=[dict(text=t, characters_including_spaces=len(t)) for t in highlights],
                    structural_coverage='All 62 expected IDs in source order, each with the five required response sections.',
                    caution='Structural completeness is not certification of completed manuscript edits or statistical validity.',
                    artifacts={f.name:digest(f) for f in [markdown,docx,args.output/'逐条覆盖核查.csv']})
    (args.output / 'document_build_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps(dict(responses=len(coverage),source_characters=len(complete),output=str(args.output),highlights=manifest['highlights']), ensure_ascii=False))


if __name__ == '__main__':
    main()
