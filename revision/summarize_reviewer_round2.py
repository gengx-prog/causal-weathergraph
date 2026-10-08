"""Build the reviewer-by-reviewer second-round experiment ledger and report."""
from pathlib import Path
from datetime import datetime,timezone
import copy
import hashlib
import json
import re
import pandas as pd
from docx import Document
from docx.shared import Inches,Mm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT.parent
OUT=BASE/'revision_outputs/reviewer_experiment_round2'
TITLE='逐条审稿意见追加实验与证据_第二轮_20260930'
LABELS={'own3':'目标历史3步','own12':'目标历史12步','own3_physical_before12':'目标3步+源前物理场',
        'wind_to_humidity':'风→湿度','humidity_to_cloud_cover':'湿度→云量','wind_to_cloud_cover':'风→云量',
        'high_minus_middle':'高状态−中状态','low_minus_middle':'低状态−中状态'}


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def read(p):return json.loads(Path(p).read_text(encoding='utf-8'))


def table(df):
    def cell(v):
        if isinstance(v,float):return f'{v:.6g}'
        return LABELS.get(str(v),str(v)).replace('|','/').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(df.columns)+' |','| '+' | '.join(['---']*len(df.columns))+' |']+
                     ['| '+' | '.join(cell(v) for v in row)+' |' for row in df.itertuples(index=False,name=None)])


def make_coverage(graph):
    original=read(OUT/'reviewer_coverage_initial.json')
    source_audit=read(OUT/'coverage_independent_audit.json')
    assert source_audit['all_56_quoted_requests_found_in_original_sources']
    sources={item['id']:item for item in source_audit['original_request_crosswalk']}
    rows=copy.deepcopy(original['items'])
    updates={
        'R1-M1':([], '依赖稳健推断已检查；校准失败保留', '不能把HAC或块聚类名义p值写成全图5%FDR保证。'),
        'R1-M3':(['graph_overlap_null_round2'],'精确约束检查与多链诊断完成','原支持退化与扩展支持混合诊断分别报告；不得把保存状态当独立随机图。'),
        'R1-M6':([], '回顾留出完成；严格未见验证不可追溯恢复','后期资料已被旧稿使用。若保留未见数据确认的主张，需要事先锁定的新资料；本轮未启动新下载。'),
        'R1-D5':(['latitude_contrasts','external_nino_physical_joint','lag_window_sensitivity','graph_overlap_null_round2'], '新增直接对比已完成', '纬度只覆盖三个环；状态和路径仍为条件预测对比，自然间接效应未识别。'),
        'R1-D6':(['lag_window_sensitivity','long_lag_benchmark'],'同样本窗口与冻结预测补充完成','后期重拟合系数与训练冻结预测分开；探索性区间没有完成总体校准。'),
        'R2-1':([], '已完成残差、HAC、块抽样及真零仿真；未通过全图校准', '保留已观察校准失败；本轮没有事后挑选块长或宣布修复。强FDR主张必须撤回或另有独立方法证据。'),
        'R2-2':(['graph_overlap_null_round2'],'保度保距离精确参照与扩展抽样完成','原支持可行图数太少；扩展支持的计算诊断不等于原支持的充分均匀随机参照。'),
        'R2-3':(['external_nino_physical_joint','latitude_contrasts'],'外部状态联合控制及纬度对比完成','外部指数不保证外生；半球季节使用当地冬夏，热带标签仍按日历定义。'),
        'R2-4':(['lag_window_sensitivity','long_lag_benchmark'],'最大源窗3/6/12及已知长lag基准完成','共同t≥15；新增物理控制用13—15步。系数重心不是传播时间，已知lag模拟不识别ERA5真图。'),
        'R2-5':(['long_lag_benchmark'],'补充已知长lag窗口覆盖基准','两变量线性DGP只补充有限lag覆盖；严格未见确认仍未满足，见R1-M6，不能把旧时期重新称独立验证。'),
        'R3-2':(['long_lag_benchmark'],'已有已知图基准外再补长lag覆盖','未包含真lag的窗口标不适用；不能把结构性排除记为算法识别失败概率。'),
        'R3-3':([], '已完成依赖诊断，校准限制继续保留','新增窗口实验沿用探索性HAC；其数值正确性不是有限样本覆盖率或FDR保证。'),
        'R3-4':(['external_nino_physical_joint'],'前置外部状态与物理控制联合完成','固定月状态广播不创造独立六小时观测；本轮不据名义显著性筛选状态结果。'),
        'R3-6':(['external_nino_physical_joint','lag_window_sensitivity'],'在前轮方向诊断基础增加联合状态和长窗控制','六物理场不能证明无遗漏驱动；前轮未来源诊断关联仍存在，方向未识别。'),
        'R3-7':(['graph_overlap_null_round2','lag_window_sensitivity'],'严格图参照诊断及效应量敏感性扩充','原支持参照退化、FDR尚未校准，不以高发现率或平均正增益证明准确。'),
    }
    for row in rows:
        extra,status,limit=updates.get(row['id'],([], '已有证据待正文/展示落实',row['current_gap']))
        row['initial_gap']=row['current_gap']
        row['current_gap']=limit
        row['initial_status']=row.get('status')
        row['status']=status
        row['primary_source_file']=sources[row['id']]['primary_source']
        row['original_quote_verified']=sources[row['id']]['normalized_quote_exact_match']
        if row['id']=='R2-5':row['requires_new_data_for_strict_claim']=True
        row['round2_experiments']=extra;row['round2_status']=status;row['remaining_limit_or_action']=limit
        row['round2_evidence']=[str(BASE/'revision_outputs'/x) for x in extra]
        row['manuscript_integration']='pending_final_manuscript_and_marked_response'
        if row['id'] in updates:row['next_action']=limit
    editor=(ROOT/'manuscript/editor_requirements.md').read_text(encoding='utf-8')
    titles=['参考文献总量与自引/期刊/作者上限','逐点单独引用、避免成组引用','首次引用顺序编号','每条highlight≤80字符','全文英文校对与修改痕迹','只引用真正相关文献']
    pending=['最终正文与文献删改后重审；不能用现有44条总数代替各期刊/作者上限检查。',
             '正文成组引用仍需逐句整理；新实验不自动完成此项。','最终编译后按首次出现顺序核查。',
             '现有五条工作稿长度71/70/61/66/71；最终版本仍需重数。',
             '代码及实验更改已追踪；全文校对、标记稿及最终逐条回复页行号仍待完成。',
             '新方法资料不自动加入论文；按最终论点逐条审核相关性及引用上限。']
    for i,(title,limit) in enumerate(zip(titles,pending),1):
        quote=re.search(rf'^{i}\. (.+)$',editor,re.M).group(1)
        rows.append(dict(id=f'EDITOR-{i}',title=title,review_request_verbatim=quote,
            source_file=str(ROOT/'manuscript/editor_requirements.md'),primary_next_action_category='文字',
            round2_experiments=[],round2_evidence=[],round2_status='编辑要求；按最终稿件验收',remaining_limit_or_action=limit,
            manuscript_integration='pending_final_manuscript_and_marked_response'))
    assert len(rows)==62 and len({r['id'] for r in rows})==62
    ledger={'updated_utc':datetime.now(timezone.utc).isoformat(),'initial_ledger_sha256':sha(OUT/'reviewer_coverage_initial.json'),
            'scope':'56 reviewer subitems plus6 editor requirements; computational completion is distinct from response/manuscript completion',
            'graph_evidence_review':graph['status'],'items':rows}
    (OUT/'reviewer_coverage_final.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2),encoding='utf-8')
    compact=pd.DataFrame([dict(意见ID=r['id'],内容=r['title'],本轮新增='、'.join(r['round2_experiments']) or '复用已有证据/文字落实',
        执行状态=r['round2_status'],未完成或解释边界=r['remaining_limit_or_action']) for r in rows])
    compact.to_csv(OUT/'reviewer_coverage_final.csv',index=False,encoding='utf-8-sig')
    return rows,compact


def main():
    graph=read(OUT/'graph_summary_for_report.json')
    for name in ('latitude_contrasts','lag_window_sensitivity','long_lag_benchmark','graph_overlap_null_round2'):
        audit=read(BASE/'revision_outputs'/name/'independent_audit.json')
        assert audit['status'].lower() in ('pass','passed','passed_with_reporting_caveat','passed_with_scope_limits'),name
    nino_audit=read(BASE/'revision_outputs/external_nino_physical_joint/implementation_audit.json')
    assert nino_audit['all_group_svd_checks_pass'] and nino_audit['all_full_design_bartlett_checks_pass']
    rows,coverage=make_coverage(graph)
    latitude=pd.read_csv(BASE/'revision_outputs/latitude_contrasts/latitude_contrasts.csv')
    lat=latitude[latitude.period.eq('evaluation')&latitude.view.eq('NH_minus_SH')&latitude.statistic.eq('linear_trend_per_10deg')&
                 latitude.edge_type.isin(['wind_to_humidity','humidity_to_cloud_cover','wind_to_cloud_cover'])].copy()
    for b in (64,128):lat[f'HAC{b}探索区间']=lat.apply(lambda r:f'[{r[f"ci95_low_hac{b}"]:.6f}, {r[f"ci95_high_hac{b}"]:.6f}]',axis=1)
    lat=lat[['edge_type','lag','candidate_count','effect','HAC64探索区间','HAC128探索区间']].rename(columns={'edge_type':'关系','lag':'源lag','candidate_count':'镜像对数','effect':'每10°有符号变化'})
    nino=pd.read_csv(BASE/'revision_outputs/external_nino_physical_joint/matched_model_comparisons.csv')
    nino=nino[nino.period.eq('evaluation')&nino.hac_bandwidth_six_hour_rows.eq(512)].groupby('contrast',as_index=False).agg(
        对比数=('pair_id','size'),原系数差中位数=('baseline_delta_beta','median'),物理控制后中位数=('physical_delta_beta','median'),符号改变数=('sign_changed','sum'))
    nino=nino.rename(columns={'contrast':'状态对比'})
    windows=pd.read_csv(BASE/'revision_outputs/lag_window_sensitivity/window_extension_summary.csv')
    windows=windows[windows.period.eq('evaluation_all')].drop(columns='period').rename(columns={
        'control_model':'控制方案','shorter_window':'原窗口','longer_window':'新窗口','models':'变量对数','longer_improves':'长窗误差下降数','median_delta_r2':'增量R方中位数'})
    centroid=pd.read_csv(BASE/'revision_outputs/lag_window_sensitivity/coefficient_centroids.csv')
    centroid=centroid[centroid.period.eq('training')].groupby(['control_model','max_source_lag']).absolute_coefficient_centroid.median().unstack().reset_index()
    centroid.columns=['控制方案','窗口3','窗口6','窗口12']
    synth=pd.read_csv(BASE/'revision_outputs/long_lag_benchmark/summary.csv')
    synth=synth[synth.target_history.eq(3)].copy()
    synth['最大系数lag恢复']=synth.apply(lambda r:f'{r.peak_correct}/100' if r.true_lag_in_window else '不适用：真lag未纳入',axis=1)
    synth=synth[['true_lag','max_source_lag','最大系数lag恢复','median_coefficient_centroid']].rename(columns={'true_lag':'真实lag','max_source_lag':'最大源lag','median_coefficient_centroid':'系数重心中位数'})
    sections=[
        '# 逐条审稿意见追加实验与证据：第二轮',
        '本轮按56条审稿细项及6条编辑要求建立执行台账，追加四组直接相关实验，并为源滞后窗口实验增加已知真lag的模拟。所有实际模型已经运行，结果另存；旧数据、实验输出和前轮报告保留。实验完成、科学主张成立、稿件文字落实是三个不同状态。本报告不是最终投稿回复。当地执行日期为2026年9月30日，日志为10月1日UTC。',
        '## 本轮完成项目',
        table(pd.DataFrame([
            ['R1-D5；R2-3','正式纬度与半球差异','756项对比；HAC64/128；三纬度环'],
            ['R2-3；R3-4/6','Niño状态×六物理控制','2592行（包含三个HAC带宽重复报告）'],
            ['R2-4；R1-D6','最大源窗3/6/12与目标历史','4536系数、7776删除损失、972窗口差'],
            ['R2-4/5；R3-2','已知lag2/6/10配对模拟','100独立重复、1800模型配置'],
            ['R2-2；R1-M3；R3-7','保度保距离图参照与重叠',graph['completion_summary']]],columns=['审稿意见','新增内容','实际产物'])),
        '## R1-D5、R2-3：正式纬度差异',
        '沿用425个镜像候选，分别拟合北半球DJF−JJA与南半球JJA−DJF的有符号条件预测斜率差。对每种关系及lag，计算三纬度带均值、带间直接差及每10°绝对目标纬度的线性趋势，并报告NH、SH和NH−SH。先在真实日历上合并跨候选、跨半球影响分数再计算HAC，未把候选边当独立样本。全部两时期×六关系×三lag×三视角×七统计，共756项，均保留两个带宽。',
        table(lat),
        '评价期风→湿度的NH−SH纬度趋势为正；湿度→云量的相应探索区间跨零，不能概括成所有关系都随纬度增强。这里只存在14.0625°、45°、75.9375°三个绝对目标纬度环，每带11个目标区域；表中三种关系的镜像pair数各为32/33/33，评价期季节窗口含DJF2528、JJA2576个六小时时次。线性趋势不代表密集纬度覆盖或纬度干预效应。采用固定候选等权估量，非全球面积加权；地理与候选组成仍混杂，热带当地冬夏只是日历标签。这一组控制目标自身历史，没有联合加入新六物理场。',
        '!FIG:latitude_trends.png',
        '## R2-3、R3-4/6：外部状态与物理控制联合',
        '复用36个冻结变量区域对及三个单独源lag。Niño3.4的训练月分位阈值、前置完整月规则不变；基线与联合模型使用相同t≥6响应。每状态具有独立截距、目标历史、控制和源斜率；六字段来自去重源/目标区域t−4..6，早于三个待检源。分组计算之后在完整日历上构造状态斜率差分数，保留跨状态协方差。',
        table(nino),
        '表中只展示评价期512步带宽对应的108项点估计对比；三个带宽的点估计相同，64/128/512全部保留在机器表。物理控制会改变部分对比的符号，因此不能把先前状态差异直接当作稳定机制。月广播不是新独立六小时观测；外部指数不等于外生处理。控制归一化条件数最高约6.30×10⁵，全部满秩仍存在强近共线性。2592行包含同一估计在三个带宽下重复报告，不是2592个独立实验。',
        '## R2-4：搜索窗口、单lag贡献与冻结预测',
        '既有实验已经在12步联合模型里删除不同lag，本轮真正新增的是最大源窗3、6、12之间的同样本比较，以及目标自身历史3步、12步、3步加源前物理场三种控制。全部从t≥15开始；物理场固定取t−13、t−14、t−15，严格早于最早待检源t−12。每个源窗口的各lag联合进入，逐个删除时重新拟合其余项，并另报删除整个源块的增量。系数表分训练与评价重拟合，预测误差全部使用训练冻结参数。',
        table(windows),
        '增加源窗口并非普遍改善：加入物理场后，6→12步仅18/36对的预测误差下降，ΔR²中位数约9×10⁻⁶。这里的窗口差是短窗MSE减长窗MSE，再除评价期目标方差；不是新增每个lag贡献的简单加总。不能按这些评价结果事后选“最好窗口”并仍称独立确认。',
        table(centroid),
        '上表是每模型用|系数|归一加权lag后，在36对之间取中位数，全部来自训练拟合。均匀权重参考为(L+1)/2，即2、3.5、6.5；它不是拟合的经验零分布。本数据的重心随窗口扩大而接近相应中点，但不同窗口也改变条件系数的估量，不能把这种接近写成证明纯粹随机或物理输送时间。全部正负lag删除结果保留。全评价10228时次；分来源评价为5884+4329，另15个来源过渡时刻仅在全评价中保留。',
        '## R2-4/5、R3-2：已知长lag的有限基准',
        '追加两变量线性过程：xₜ=0.75xₜ₋₁+εₓ，yₜ=0.6yₜ₋₁+0.25xₜ₋d+εᵧ，d固定为2、6、10。100组独立随机种子在三种真lag之间共享创新，以作配对比较；每组4096训练时次、2048评价时次，训练拟合标准化。统一t≥12后训练模型行4084，评价2048。两种目标历史和三种源窗全部保留；下表及随后图(b)展示目标历史3步，历史12步结果完整保留在机器表。',
        table(synth),
        '当真lag在窗口内，各配置在这一DGP下均为100/100最大绝对训练系数lag匹配，Wilson95蒙特卡洛区间约[0.9630,1]；这不等于真实恢复概率必为1，也不覆盖其他效应强度、非线性或未观测共同原因。真lag不在窗口的条目标“不适用”，不能把其结构性零匹配率解释为算法识别失败概率。即使真lag=2始终可恢复，窗口12的系数重心仍明显高于2，进一步表明平均lag不能自动解释成传播时间。',
        '!FIG:source_window_centroids.png',
        '## R2-2、R1-M3、R3-7：图参照、约束退化与混合',
        *graph['paragraphs'],
        table(pd.DataFrame(graph['table_rows'])),
        '## 数值实现与原文覆盖核查',
        '纬度实验独立重算全部756项点估计，并核验262个首阶段SVD拟合与42组日历HAC聚合；状态联合模型核验72个分组SVD及24个完整设计/Bartlett计算；窗口实验核验24个联合拟合、288个删除损失及36个窗口差。长lag基准独立重生成三组指定种子的过程并核验18个模型，全部汇总表另行复算。图参照独立检查全部1600个保存图、原支持全部可行图及19项哈希，并重放539440次提案；这些检查通过。对原始Reviewer 1附件及Reviewer 2/3评语逐项核对，56个答复细项的原文均找到对应；其中Reviewer 1的44项是内部拆分编号，并非原文44个独立问题。这些是实现与覆盖检查，不能证明全部科学假设或投稿要求已经满足。',
        '## 仍需明确保留的技术边界',
        '依赖稳健检验与真零模拟已按审稿要求执行，但结果没有建立全图5%FDR保证。新增实验没有改变这一事实，也没有挑选有利的HAC带宽。若论文坚持全图错误率控制，需要新的有效推断方案和独立校准；就现有证据，应收窄为条件预测关系和探索性敏感性。',
        '真实大气因果图、自然间接效应、全层水汽预算和真实运动轨迹不由现有区域回归识别。前轮方向/水汽通量代理已完成，但完整物理预算需要更充分的垂直资料、边界项和识别假设。52/60路径缺合格距离替代中介的问题继续保留。资源限制是本次取舍原因，不是科学上永久不可实现。',
        '历史上已经使用过的2019—2025年不能通过重新切分变成从未查看的独立确认集。当前是锁定流程后的回顾时间留出；若坚持前瞻确认，需要事先冻结的新资料。CERES目前仍仅2019年1月两个区域，未新增多年下载。',
        '## 逐条执行台账：56条审稿细项与6条编辑要求',
        '以下逐条对应审稿原文拆分后的细项，不把每个文字要求改写成新实验。完整原文、原文件位置、既有证据、新增目录、剩余事项及稿件状态保存在reviewer_coverage_final.json；同内容的可筛选表为reviewer_coverage_final.csv。编辑要求仍按最终正文、参考文献及标记稿验收。',
        table(coverage[['意见ID','内容','执行状态','未完成或解释边界']]),
        '## 复核与复现入口',
        'revision_outputs/latitude_contrasts/：756项对比、2550条首阶段系数、日历分数、独立SVD与Bartlett核验。\n\nrevision_outputs/external_nino_physical_joint/：2592行联合对比、同样本变化、原生月/episode数和共线性诊断。\n\nrevision_outputs/lag_window_sensitivity/：全部联合系数、训练冻结删除损失、窗口差、重心和独立审计。\n\nrevision_outputs/long_lag_benchmark/：100配对重复的真实参数、所有系数、预测误差及独立重现。\n\nrevision_outputs/graph_overlap_null_round2/：原支持精确枚举、扩展支持多链、约束和混合诊断。\n\nrevision_outputs/reviewer_experiment_round2/：62细项台账、来源独立核查、绘图数据及报告验收记录。',
        '代码和复现命令保存在causal-weathergraph/revision/；每批次包含执行时的代码快照与输入输出哈希。数值复核通过只说明实现与独立计算一致，不证明物理识别或统计覆盖成立。正文、高亮稿和最终英文逐条回复仍需准确纳入本轮证据。',
    ]
    rendered=[]
    for section in sections:
        if section.startswith('!FIG:'):
            name=section.split(':',1)[1]
            rendered.append(f'![追加实验图](revision_outputs/reviewer_experiment_round2/figures/{name})')
        else:rendered.append(section)
    markdown='\n\n'.join(rendered)+'\n'
    (BASE/(TITLE+'.md')).write_text(markdown,encoding='utf-8')
    (ROOT/'revision/REVIEWER_EXPERIMENTS_ROUND2_ZH.md').write_text(markdown.replace('(revision_outputs/','(../../revision_outputs/'),encoding='utf-8')
    doc=Document();page=doc.sections[0]
    page.page_width,page.page_height=Mm(210),Mm(297)
    page.top_margin=page.bottom_margin=Inches(.65);page.left_margin=page.right_margin=Inches(.65)
    style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(20);doc.styles['Heading 1'].font.size=Pt(14)
    for section in sections:
        if section.startswith('# '):doc.add_heading(section[2:],0)
        elif section.startswith('## '):doc.add_heading(section[3:],1)
        elif section.startswith('!FIG:'):doc.add_picture(str(OUT/'figures'/section.split(':',1)[1]),width=Inches(6.75))
        elif section.startswith('| '):
            data=[r.strip('| ').split(' | ') for r in section.splitlines() if not r.startswith('| ---')]
            tab=doc.add_table(rows=1,cols=len(data[0]));tab.style='Light Shading Accent 1'
            tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for c,v in zip(tab.rows[0].cells,data[0]):c.text=v
            for values in data[1:]:
                for c,v in zip(tab.add_row().cells,values):c.text=v
            for row in tab.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1
                        for run in para.runs:run.font.size=Pt(8.5)
        else:
            for para in section.split('\n\n'):doc.add_paragraph(para)
    footer=page.footer.paragraphs[0];footer.text='2026-09-30 · 审稿追加实验第二轮 · '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(BASE/(TITLE+'.docx'))
    save={'generated_utc':datetime.now(timezone.utc).isoformat(),'coverage_items':62,'table_count':len(doc.tables),
          'figure_count':len(doc.inline_shapes),'report_files':{ext:sha(BASE/(TITLE+ext)) for ext in ('.md','.docx')},
          'script_sha256':sha(__file__),'graph_review_sha256':sha(OUT/'graph_summary_for_report.json')}
    (OUT/'report_manifest.json').write_text(json.dumps(save,indent=2),encoding='utf-8')
    print(json.dumps(save,indent=2))


if __name__=='__main__':main()
