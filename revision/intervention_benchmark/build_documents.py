"""Scientific supplement in the project's existing local document format."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT.parent / "revision_outputs/intervention_benchmark_v1"
METHODS = ["unadjusted_ols", "local_ols", "observed_z_ols", "local_hgb_gcomp", "observed_z_hgb_gcomp", "full_hgb_gcomp", "oracle_ipw", "oracle_aipw"]
LABELS = ["未调整 OLS", "局地控制 OLS", "加入 Z 的 OLS", "局地 HGB", "加入 Z 的 HGB", "全状态 HGB", "已知倾向 IPW", "已知倾向 AIPW"]
EN_LABELS = ["Unadjusted OLS", "Local OLS", "Observed-Z OLS", "Local HGB", "Observed-Z HGB", "Full-state HGB", "Oracle IPW", "Oracle AIPW"]
SCENARIOS = ["randomized_transport", "confounded_transport", "confounded_feedback", "confounded_null"]
SC_LABELS = ["随机干预与输送", "共同驱动与输送", "共同驱动加反馈", "共同驱动且零作用"]
EN_SCENARIOS = ["Randomized transport", "Confounded transport", "Confounded feedback", "Confounded exact null"]


def sha(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def table(headers, rows):
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"]*len(headers)) + " |"] + ["| " + " | ".join(str(v) for v in row) + " |" for row in rows])


def main():
    audit = json.loads((OUT / "independent_audit.json").read_text(encoding="utf8"))
    assert audit["status"] == "PASS_NUMERICS_MODELS_METRICS_SUMMARIES"
    manifest = json.loads((OUT / "run_manifest.json").read_text(encoding="utf8"))
    assert manifest["status"] == "completed" and len(manifest["completed"]) == 80
    design = json.loads((OUT / "design.json").read_text(encoding="utf8"))
    data = pd.read_csv(OUT / "summary.csv")
    metrics = pd.read_csv(OUT / "all_metrics.csv")
    resolution = pd.read_csv(OUT / "resolution_sensitivity.csv")
    def row(sc, method, h=6, target="cloud_proxy_pp"):
        s = data[(data.scenario == sc) & (data.method == method) & (data.horizon_hours == h) & (data.outcome == target)]
        assert len(s) == 1
        return s.iloc[0]
    def f(value): return f"{value:.4f}"
    truth_rows = []
    for sc, label in zip(SCENARIOS, SC_LABELS):
        truth_rows.append([label] + [f(row(sc,"unadjusted_ols",h,target).mean_test_bank_truth_mean) for target,h in [("vapor_g_per_kg",6),("vapor_g_per_kg",12),("cloud_proxy_pp",6),("cloud_proxy_pp",12)]])
    comparison_rows = []
    for method,label in zip(METHODS,LABELS):
        r = row("confounded_transport",method)
        comparison_rows.append([label,f(r.mean_effect_estimate),f(r.mean_signed_error_vs_test_bank),f(r.ate_rmse),f(r.mean_factual_test_mse) if np.isfinite(r.get("mean_factual_test_mse",np.nan)) else "不适用"])
    null_rows = []
    for method,label in zip(METHODS,LABELS):
        r = row("confounded_null",method)
        null_rows.append([label,f(r.mean_effect_estimate),f(r.ate_rmse),f"{int(r.null_rejections)}/{int(r.null_tests)}" if np.isfinite(r.get("null_tests",np.nan)) else "未构造区间", f"[{r.null_wilson_low:.2f}, {r.null_wilson_high:.2f}]" if np.isfinite(r.get("null_tests",np.nan)) else "—"])
    plt.rcParams.update({"font.size":10,"axes.spines.top":False,"axes.spines.right":False})
    fig, axes_grid = plt.subplots(2,2,figsize=(8.5,7.6),sharey=True)
    axes = axes_grid.ravel()
    plot_rows = []
    for ax,sc,title in zip(axes,SCENARIOS,EN_SCENARIOS):
        for i,method in enumerate(METHODS):
            r = row(sc,method)
            value = r.mean_signed_error_vs_test_bank
            lo,hi = r.mc95_low_signed_error_vs_test_bank,r.mc95_high_signed_error_vs_test_bank
            ax.errorbar(value,7-i,xerr=[[max(0,value-lo)],[max(0,hi-value)]],fmt="o",color="#287C8E" if "oracle" not in method else "#A54C3B",capsize=3,markersize=4)
            plot_rows.append({"scenario":sc,"method":method,"target":"cloud_proxy_pp","horizon_hours":6,"mean_bias":value,"mc95_low":lo,"mc95_high":hi})
        ax.axvline(0,color="0.4",linestyle="--",lw=1)
        ax.set_title(title,fontsize=10); ax.xaxis.set_major_locator(MaxNLocator(4))
        ax.grid(axis="x",alpha=.15); ax.set_xlabel("Mean effect error (proxy pp)")
    axes[0].set_yticks(range(8),EN_LABELS[::-1])
    axes[2].set_yticks(range(8),EN_LABELS[::-1])
    fig.suptitle("Six-hour cloud-proxy effect recovery",fontsize=13)
    fig.text(.5,.015,"Intervals: Monte Carlo uncertainty of mean error across 20 repetitions.\nPanel x-axis scales differ.",ha="center",fontsize=9)
    fig.tight_layout(rect=[0,.06,1,.95])
    fig.savefig(OUT / "effect_recovery.png",dpi=180)
    fig.savefig(OUT / "effect_recovery.pdf")
    plt.close(fig)
    pd.DataFrame(plot_rows).to_csv(OUT / "figure_values.csv",index=False)
    null_ols = row("confounded_null","unadjusted_ols")
    null_aipw = row("confounded_null","oracle_aipw")
    conf_ols = row("confounded_transport","unadjusted_ols")
    conf_adj = row("confounded_transport","observed_z_ols")
    conf_aipw = row("confounded_transport","oracle_aipw")
    text = rf'''# 干预真值模拟实验与因果识别边界

本实验用于检验：在明确的生成方程和风场干预下，原有线性估计方法的适配版本及新增对照方法，能否恢复模拟器内部的干预效应。已完成4种情形、每种20次独立重复、每次3072个独立初始状态的模拟，并将预测误差、干预效应误差和零作用误报分别报告。每个初始状态对应两条共享外生扰动的潜在轨迹；不同情形共享初始状态和扰动，因此不能把80个情形重复单元视为相互独立。

在存在共同驱动的输送情形中，6小时云代理的配对平均效应为 {conf_ols.mean_test_bank_truth_mean:.4f} 个代理百分点；未调整OLS估计为 {conf_ols.mean_effect_estimate:.4f}，加入真实共同驱动Z的OLS为 {conf_adj.mean_effect_estimate:.4f}，已知倾向概率AIPW为 {conf_aipw.mean_effect_estimate:.4f}。这些数值评估的是指定模型中的恢复能力，不是现实大气的云量效应。

AIPW的平均估计接近配对真值，但20次重复的ATE RMSE仍为 {conf_aipw.ate_rmse:.4f} 个代理百分点，不能仅凭平均偏差小就称其稳定准确。所有方法的误差、预测性能与不利结果均在后文并列保留。

## 一 模型与干预定义

计算域为1200 km周期一维环，48格、每格25 km，积分步长300 s。状态为水汽混合比 $r_v$、凝结水混合比 $r_l$（均为每千克干空气的水质量）以及外部热力强迫下的温度T。它们不能与ERA5比湿或总云量直接等同。固定局地区间200–400 km作为目标，避免周期平移下全域平均对风不敏感的问题。

每个初始状态先生成共同驱动 $Z_0$，随后分配二元处理A。A=1在前1小时均匀叠加+2 m/s，A=0叠加−2 m/s；两种政策之差为4 m/s。一小时后撤除脉冲，允许反馈继续演化。给定相同初始状态和完全相同的未来外生噪声，分别积分两条轨迹：

$$
\tau_i(h)=Y_i(h;A=1)-Y_i(h;A=0),\qquad
\bar\tau(h)=\frac{{1}}{{N}}\sum_i\tau_i(h).
$$

这是两种风场政策的模型内总效应，包含温度、凝结水输送及反馈的响应，不能直接称为“仅经湿度传递”的自然间接效应。原始模型内效应也含数值离散误差，下面另外报告分辨率敏感性。

风由外部过程与可选反馈给定：$U=8+2\tanh Z+1.5\tanh\eta+2b\tanh(\bar r_l/0.001-1)+I_{{t<1h}}(4A-2)$，单位m/s；b在反馈情形取1，其余为0。Z和η为独立OU外生过程，相关时间分别为6 h与3 h。模拟不解动量、压力或完整Navier–Stokes方程；该反馈仅为人为规定的识别压力测试。

水汽、凝结水与温度采用周期守恒迎风输送。水汽向空间非均匀参考场和Z控制的湿源松弛，时间尺度12 h；温度向空间非均匀参考场和Z控制的热力环境松弛，时间尺度6 h。在每个步长中，固定当前T执行饱和调整：

$$
D=\max(r_v-r_s(T),0)-\min\{{r_l,\max(r_s(T)-r_v,0)\}},\quad
r_v\leftarrow r_v-D,\quad r_l\leftarrow r_l+D.
$$

再以6 h时间尺度移除凝结水并计入降水汇。所有水汽松弛源汇与凝结水移除均进入水量台账。此模型没有潜热回馈，因而没有声称能量闭合或完整云微物理。参考文献中的湿模式提供建模背景，本实现不等同于Hydro-ABC等已发表模式。

云代理定义为 $C_*=100\,\operatorname{{mean}}_R[1-\exp(-r_l/0.001)]$。先逐格非线性转换再区域平均；它是0–100范围的人工凝结水诊断，文中“代理百分点”仅指该量的差值。

## 二 识别条件与四种对照

随机情形采用 $P(A=1)=0.5$；其余情形采用 $e(Z_0)=\operatorname{{clip}}[\operatorname{{expit}}(1.5Z_0),0.15,0.85]$。Z同时影响处理分配和后续热湿强迫，从而产生共同驱动混杂。由于分配机制只使用干预前Z，在该生成器中观测Z足以阻断处理的后门路径；无需假定观测了大气全部状态。局地特征包含风这一Z的代理，因此“未直接观测Z”不代表完全没有关于Z的信息。

| 情形 | 干预分配 | 风进入状态方程 | 反馈 |
| --- | --- | --- | --- |
| 随机干预与输送 | 随机 | 水汽、凝结水、温度平流 | 无 |
| 共同驱动与输送 | 由Z决定倾向 | 同上 | 无 |
| 共同驱动加反馈 | 由Z决定倾向 | 同上 | 凝结水反馈风 |
| 共同驱动且零作用 | 由Z决定倾向 | 全部关闭 | 无 |

零作用情形仅用于检验混杂造成的误报，不是现实大气假设。该情形在所有格点、所有时间步的两臂状态哈希完全一致，且每个样本的干预效应严格为零；仍允许A与结果在观测上相关。

## 三 拟合与评价

每次重复固定划分2048训练、512验证、512测试初始状态。训练仅使用被分配的一个处理臂和相应事实结果；配对真值只在测试集用于评价。固定超参数，不使用验证集或测试效应挑选有利模型。两种结局为6 h和12 h的局地水汽混合比（g/kg干空气）及云代理。

线性方法复用原 `fit_edge` 的FWL/OLS系数，改成仅调整干预前状态的终点效应估计。它是明确标注的适配版本，不是直接验证原时间滞后网络或旧显著性加权评分。新增HGB拟合事实条件均值，再计算同一初始状态下A=1和A=0的预测差。局地控制、局地加Z、全格点加Z是信息充分性比较，不能称为同信息模型排名。

IPW使用已知生成倾向；AIPW使用该倾向及独立训练集拟合的“局地加Z”HGB。它们是已知机制的oracle对照，不能把这里已知的倾向概率当作真实大气中已得到的概率。OLS区间为HC3线性投影系数的渐近区间；oracle区间为独立样本影响分数的正态区间。HGB未生成单次拟合的效应区间。所有效应误差在20次重复上汇总，重复均值的t区间只描述Monte Carlo误差。有限测试样本的配对真值本身也有抽样波动，不把它当作无误差的总体效应。

## 四 配对干预效应

下表为每次测试集配对效应均值在20次重复上的平均。正负号由实际模拟输出决定，没有预先规定“更大风速必然增加湿度或云”。

{table(['情形','水汽6 h','水汽12 h','云代理6 h','云代理12 h'],truth_rows)}

水汽单位g/kg干空气；云单位代理百分点。随机与混杂输送情形使用相同潜在轨迹，区别只在事实处理分配，因此这两行情形的真值相同是设计结果。

## 五 效应恢复与预测的区别

共同驱动与输送情形，6 h云代理的全部方法结果如下。ATE RMSE为20次“估计平均效应减测试配对平均效应”的均方根；事实MSE评价预测，量纲为代理百分点平方。

{table(['方法','估计效应','平均误差','ATE RMSE','事实MSE'],comparison_rows)}

即使加入了生成器中足够的共同驱动Z，有限样本、回归形式近似和测试集抽样仍可能留下效应误差；存在效应异质性时，单一线性处理系数也未必等于目标人群的平均效应。这类误差不能一概归于识别条件不足。已知倾向IPW也可能有较大方差；应结合重复实验及其区间判断，不能依据一次估计偏离就否定其识别依据。

![四种情形的6小时云代理效应误差](effect_recovery.png)

图中区间为20次独立重复的平均误差Monte Carlo 95%区间，各面板横轴范围不同。完整8种方法、4种情形、2种结局与2个时距均保存在结果CSV，没有按改善与否筛选。

## 六 零作用对照

下面情形的真实效应逐样本为零。6 h云代理上，未调整OLS的平均估计为 {null_ols.mean_effect_estimate:.4f}，oracle AIPW为 {null_aipw.mean_effect_estimate:.4f}。

{table(['方法','平均估计','ATE RMSE','排除零的次数','次数比例Wilson区间'],null_rows)}

这里只统计已有95%区间是否排除零，不为没有区间的HGB强加显著性判断。在混杂存在时，OLS的线性投影系数可以真实不为零，因此该表描述“把关联误解为作用”的风险，不能说OLS对自身投影参数的统计检验必然失效。20次重复不足以精确校准5%误报率，更不能声称全图FDR已被校准。

## 七 数值核验及边界

守恒输送、饱和相变、非负性、CFL及逐步水源汇台账均已核验。独立审计另从保存的外生路径及状态重建受检轨迹，并重算效应与估计指标；详细覆盖范围见独立审计回执，不把抽样轨迹检查称为全部轨迹独立重算。

另固定128个初始状态，比较48格/300 s、48格/150 s和96格/150 s，保持初始连续场参数与外生OU路径一致。所有情形与结局均保留于 `resolution_sensitivity.csv`。与基线相比，加密网格后平均水汽效应最大绝对变化为 {resolution[(resolution.resolution=='refined_grid')&(resolution.outcome=='vapor_g_per_kg')].mean_effect_change.abs().max():.6f} g/kg；平均云代理效应最大绝对变化为 {resolution[(resolution.resolution=='refined_grid')&(resolution.outcome=='cloud_proxy_pp')].mean_effect_change.abs().max():.6f} 代理百分点。这是有限分辨率敏感性，不能当作连续方程极限已经严格收敛。

本实验可支持“在明示方程、干预及观测条件下评估因果效应恢复”。它不能证明ERA5风—湿度—云链已经识别，也没有识别湿度的自然中介效应。真实资料部分仍须保留共同驱动、时空聚合和反向诊断的限制。正文整合、最终参考文献总量检查以及全篇修订标记尚未完成。

## 八 可用于审稿回复的英文说明

> We added a controlled intervention benchmark with explicitly paired potential trajectories. A spatially uniform wind perturbation of +2 versus −2 m/s is applied for the first hour, with identical initial states and future exogenous innovations in both worlds. The benchmark includes randomized assignment, a measured or omitted common driver, a prescribed condensate-to-wind feedback, and an exact structural null in which wind is removed from all prognostic transport terms. The model is an idealized, externally thermostatted moist-tracer system; its condensate-based diagnostic is not ERA5 total cloud cover.

> We evaluate an explicitly adapted version of the original linear estimator, nonlinear g-computation, and known-propensity oracle comparators using only factual training outcomes. Paired intervention outcomes are withheld for evaluation. Effect-recovery errors and factual prediction errors are reported separately, including unfavorable results and exact-null diagnostics. This experiment tests recovery of specified model-internal total effects; it does not identify humidity-mediated effects or resolve the observational identification assumptions for the original atmospheric graph.

## 九 证据与相关研究

冻结设计为 `design.json`；逐样本事实结果与配对真值在 `data/`；模型与预测在 `models/`；完整指标为 `all_metrics.csv`（{len(metrics)}行），汇总为 `summary.csv`（{len(data)}行）。独立核验为 `independent_audit.json`。所有模拟来自本地生成，不需要新增大规模气象下载。

1. 水汽平流凝结的理想化建模背景：[Sukhatme and Young 的平流凝结模型研究](https://arxiv.org/abs/1105.0470)。本文生成器和参数为本次明确设定，不声称复现该论文。
2. 更完整湿动力模型的对照背景：[Zhu and Bannister 的 Hydro-ABC 模型](https://doi.org/10.5194/gmd-16-6067-2023)。该模式包含的动力和能量处理不能作为本简化模型已实现的功能。
3. 时间序列因果推断的假设与挑战：[Runge 等的地球系统因果推断研究](https://doi.org/10.1038/s41467-019-10105-3)。
4. 中介效应的额外识别条件：[Imai 等的中介识别及敏感性分析](https://doi.org/10.1214/10-STS321)。
'''
    md = OUT / "干预真值模拟实验与因果识别边界.md"
    md.write_text(text,encoding="utf8")
    sys.path.insert(0,str(ROOT))
    from revision.build_response_document import create_reference
    reference = OUT / "document_style.docx"; create_reference(reference)
    docx = md.with_suffix(".docx")
    subprocess.run([str(Path.home()/"AppData/Local/Pandoc/pandoc.exe"),str(md),"-o",str(docx),"--from=markdown+tex_math_dollars","--standalone",f"--reference-doc={reference}",f"--resource-path={OUT}"],check=True)
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    doc = Document(docx)
    doc.core_properties.title = "干预真值模拟实验与因果识别边界"
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"): p.paragraph_format.keep_with_next=True
        if p._p.xpath('.//w:drawing'): p.paragraph_format.keep_with_next=True
    for shape in doc.inline_shapes:
        shape.width = int(shape.width * .93)
        shape.height = int(shape.height * .93)
    for tab in doc.tables:
        for i, rr in enumerate(tab.rows):
            rr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in rr.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    if i<len(tab.rows)-1:p.paragraph_format.keep_with_next=True
                    for run in p.runs:run.font.size=Pt(9)
    doc.save(docx)
    delivery = dict(created_utc=datetime.now(timezone.utc).isoformat(), audit_sha256=sha(OUT/"independent_audit.json"), design_sha256=sha(OUT/"design.json"), source_sha256=sha(__file__), files={p.name:sha(p) for p in [md,docx,OUT/"effect_recovery.png",OUT/"figure_values.csv"]}, pdf_status="pending_export")
    (OUT / "document_manifest.json").write_text(json.dumps(delivery,ensure_ascii=False,indent=2),encoding="utf8")
    print(json.dumps(dict(docx=str(docx),rows=len(metrics)),ensure_ascii=False))


if __name__ == "__main__":
    main()
