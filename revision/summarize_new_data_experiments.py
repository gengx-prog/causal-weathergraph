"""Generate a compact evidence-backed Chinese report from completed outputs."""
from pathlib import Path
from datetime import datetime, timezone
import json
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
MAJOR = ROOT.parent
OUT = MAJOR / 'revision_outputs'
LABELS = {
    'own_history': '自身历史',
    'physical_response_history': '响应前物理历史',
    'physical_presource_history': '源前物理历史',
    'baseline': '原路径基线',
    'physical_before_path_source': '最早源前物理历史',
    'humidity_to_cloud_cover': '湿度→云量',
    'wind_to_cloud_cover': '风→云量',
    'wind_to_humidity': '风→湿度',
    'original': '原方向', 'reversed': '反方向',
    'past': '过去源', 'future_diagnostic': '未来源诊断',
    'evaluation_all': '2019—2025全部',
    'evaluation_wb2': '早期来源评价段', 'evaluation_cds': 'CDS评价段',
    'ordinary': '普通', 'strong': '强', 'unselected': '未选中',
    'high_minus_middle': '高状态−中状态', 'low_minus_middle': '低状态−中状态',
    'north_atlantic': '北大西洋', 'south_atlantic': '南大西洋',
    'ceres_24_slots_vs_era5_4_times': '24槽对4整点',
    'ceres_4_halfhour_slots_vs_era5_4_times': '4半点槽对4整点',
}


def table(frame):
    def cell(value):
        return f'{value:.6g}' if isinstance(value, float) else LABELS.get(str(value), str(value))
    lines = ['| ' + ' | '.join(frame.columns) + ' |', '| ' + ' | '.join(['---'] * len(frame.columns)) + ' |']
    lines += ['| ' + ' | '.join(cell(v) for v in row) + ' |' for row in frame.itertuples(index=False, name=None)]
    return '\n'.join(lines)


def main():
    directory = OUT / 'physical_controls_experiments'
    run = json.loads((directory / 'manifest.json').read_text(encoding='utf-8'))
    if run['status'] != 'completed':
        raise ValueError('Model batch is not complete')
    audit = json.loads((directory / 'independent_audit.json').read_text(encoding='utf-8'))
    if audit['status'] != 'passed':
        raise ValueError('Independent physical-model audit has not passed')
    direction = pd.read_csv(directory / 'direction_frozen_losses.csv')
    paths = pd.read_csv(directory / 'path_frozen_losses.csv')
    control_compare = pd.read_csv(directory / 'direction_control_comparisons.csv')
    subset = direction[(direction.period == 'evaluation_all') & (direction.orientation == 'original') & (direction.source_timing == 'past')]
    prediction_summary = subset.groupby(['model', 'original_edge_type'], as_index=False).agg(
        模型数=('valid', 'size'), 有效数=('valid', 'sum'), 源增量为正=('mse_gain', lambda v: int((v > 0).sum())),
        增量R方中位数=('delta_r2_test_centered', 'median'))
    prediction_summary = prediction_summary.rename(columns={'model': '控制方案', 'original_edge_type': '原关系类型'})
    comparisons = control_compare[(control_compare.period == 'evaluation_all') & (control_compare.orientation == 'original') & (control_compare.source_timing == 'past')]
    whole_summary = comparisons.groupby(['model', 'original_edge_type'], as_index=False).agg(
        模型数=('whole_model_mse_improvement', 'size'), 整模型误差下降数=('whole_model_mse_improvement', lambda v: int((v > 0).sum())),
        源增量变化中位数=('source_increment_change', 'median'))
    whole_summary = whole_summary.rename(columns={'model': '控制方案', 'original_edge_type': '原关系类型'})
    diagnostic_summary = direction[direction.period == 'evaluation_all'].groupby(
        ['orientation', 'source_timing', 'model'], as_index=False).agg(
        模型数=('valid', 'size'), 源增量为正=('mse_gain', lambda v: int((v > 0).sum())),
        增量R方中位数=('delta_r2_test_centered', 'median'))
    diagnostic_summary = diagnostic_summary.rename(columns={'orientation': '方向', 'source_timing': '源时间', 'model': '控制方案'})
    segment_summary = direction[(direction.orientation == 'original') & (direction.source_timing == 'past')].groupby(
        ['period', 'model'], as_index=False).agg(
        模型数=('valid', 'size'), 源增量为正=('mse_gain', lambda v: int((v > 0).sum())),
        增量R方中位数=('delta_r2_test_centered', 'median'))
    segment_summary = segment_summary.rename(columns={'period': '评价范围', 'model': '控制方案'})
    path_subset = paths[(paths.period == 'evaluation_all') & (paths.stage == 'b_humidity_to_cloud_given_wind')]
    path_summary = path_subset.groupby(['model', 'stratum'], as_index=False).agg(
        路径数=('valid', 'size'), 湿度增量为正=('mse_gain', lambda v: int((v > 0).sum())),
        增量R方中位数=('delta_r2_test_centered', 'median'))
    path_summary = path_summary.rename(columns={'model': '控制方案', 'stratum': '原训练分层'})
    ceres = pd.read_csv(OUT / 'ceres_pilot_comparison/descriptive_metrics.csv')
    ceres_table = ceres[['region', 'scheme', 'n_paired_days', 'bias_pp_ceres_minus_era5', 'rmse_pp', 'pearson_r']].rename(columns={
        'region': '区域', 'scheme': '采样方案', 'n_paired_days': '配对日数', 'bias_pp_ceres_minus_era5': '偏差(百分点)', 'rmse_pp': 'RMSE(百分点)', 'pearson_r': '相关系数'})
    nino_dir = OUT / 'external_nino_regimes'
    nino = pd.read_csv(nino_dir / 'nino_regime_contrasts.csv')
    nino_summary = nino[(nino.period == 'evaluation') & (nino.hac_bandwidth_six_hour_rows == 512)].groupby(['contrast', 'edge_type'], as_index=False).agg(
        模型对比数=('delta_beta', 'size'), 有符号斜率差中位数=('delta_beta', 'median'), 探索性标准误中位数=('se_hac', 'median'))
    nino_summary = nino_summary.rename(columns={'contrast': '状态对比', 'edge_type': '关系类型'})
    sections = [
        '# 新下载数据处理与补充实验结果（2026-09-30）',
        '本报告记录实际执行的新增工作。所有结果保留源文件、冻结设计、代码与输入输出哈希；原有实验和原始数据未覆盖。下列结果是修订阶段的补充敏感性分析，不能写成全部审稿问题已经解决。',
        '## 数据处理',
        '144份CDS原始文件已完成全值检查、SHA256核验和球面面积保守重网格，派生网格与早期WeatherBench2控制场一致；压缩派生NetCDF共148,884,431 bytes。独立审计另行核验全部派生值、时间、坐标、单位与哈希。',
        '六字段为500/700 hPa垂直气压速度、500 hPa位势、700 hPa温度、地面气压和海平面气压。1979—2025年共68,668个连续六小时时次、66个区域；58,440个训练时次与10,228个评价时次按原设计分开。每格月气候态、异常中心和尺度仅在1979—2018拟合，再对原区域内标准化格点做算术均值。原物理区域均值单独保存，不能混为相同估计量。',
        '已检查评价值改成NaN或极大值不会改变训练参数，并进行原NetCDF直接复算。另存地面气压低于500/700/850 hPa的粗网格提示；这是地下压力层适用性诊断，不能替代原生地形掩码，也没有据此插补或删除原值。',
        '## 新增共同驱动控制与方向诊断',
        '复用既有训练期冻结的36个变量区域对，保留原/反两个方向和1—3步过去源、1—3步未来源；三套控制均使用共同响应窗口。baseline为目标自身1—3步历史；response_history增加源/目标地区六物理字段1—3步历史，可能晚于被测过去源；presource_history改为4—6步历史，严格早于所有被测过去源。未来信息重构只用于反时序诊断，不计为可实施预测。',
        '所有预测参数在训练期冻结。训练要求未来诊断源也不跨入2019；整评价期保留来源衔接行，分来源表则剔除任何输入跨2023-01-11边界的设计行。评价期重新拟合的系数表与冻结预测误差表明确分开。新增变量是作者按物理动机选择的实现，不能写成审稿人逐项指定了六个字段。',
        '以下是原方向、过去源在2019—2025评价期的结果，每套控制方案均有108个模型（36对×3个单独源滞后）。ΔR²=(MSE受限−MSE完整)/评价期目标均值中心化方差；正源增量表示同一控制集内加入该源改善预测。正增量计数不代表显著、独立或因果关系。',
        table(prediction_summary),
        '整模型改善与源增量变化是不同问题。下表“源增量变化”为控制后与基线的(MSE受限−MSE完整)之差，单位是区域标准化响应的平方，不是ΔR²之差。整模型误差下降按MSE判断，保留所有负面结果：',
        table(whole_summary),
        '源前物理历史方案有105/108个完整模型误差下降，响应前物理历史方案为108/108；但源自身增量提高的模型分别只有52/108和46/108。物理场提供额外预测信息，并不意味着原有每个源的独立预测贡献都增强。',
        '反方向和未来源诊断结果如下。未来源也经常改善重构，物理控制并未消除反时序关联；这些模式不支持仅靠预测增量判断物理方向。不同方向的响应变量不同，不能把跨方向R²简单当作同一物理效应的优劣排序。',
        table(diagnostic_summary),
        '来源分段敏感性如下：早期来源评价段为2019-01-01至2023-01-10，CDS段从2023-01-11开始。分段剔除跨边界输入行，日期端点还受共同未来诊断窗口约束；具体样本数逐行保存在结果表。时间段、气候背景及原主变量处理方法同时变化，段间差异不能归因于单一数据来源。',
        table(segment_summary),
        '## 时间对齐路径补充',
        '复用全部60条冻结路径，始终按W_i(t−l1−l2)→H_h(t−l2)→C_j(t)对齐。新增六字段来自去重后的i/h/j地区，时间严格在最早W之前1—3步。与原H/T、C/T历史基线作同样本比较。中介首段的评价拟合排除仍属于2018年的H响应，故a与b的起始样本略有不同；即使支持完全一致，本分析也未建立中介识别假设，a×b只能是描述性乘积，不能视为自然间接效应。原来仅8/60存在合格距离替代中介的问题并未解决。',
        table(path_summary),
        '强路径组的湿度增量ΔR²中位数从0.00197678降至0.000703600；20条中正增量由20条变为18条。该固定样本显示路径解释对物理控制敏感，不能据此宣称湿度中介机制已经证明；分层间也没有进行随机化效应比较。',
        '## 外部Niño3.4状态',
        '使用已下载月异常指数定义训练分位状态，480个原生训练月份的q20/q80为−0.66/+0.66°C。每个响应使用“响应前24小时所在月的前一个完整月”指数，确保指数月份结束早于被测1—3步源；开头缺少1978年11月的行保留为缺失并排除。月指数广播到六小时行并不创造新独立观测；原生月份数、连续状态段与模型行数分别记录。',
        '复用36个冻结变量区域对的原方向与三个单独源滞后。这是一组独立的状态敏感性实验，控制目标自身1—3步历史，允许各状态的截距、历史系数和源斜率不同；没有联合加入上节的六个新物理控制，不能宣称已经完成“外部状态×物理控制”的联合稳健性检验。',
        '高减中、低减中采用状态特异截距、目标历史与source×state直接交互，训练/评价分别回顾拟合。HAC64/128/512三个带宽（16/32/128天）全部保留，无BH或按显著性选择方案。完整1,296行包含同一系数对比在三个带宽下的重复报告，不能称为1,296个独立检验。下表展示评价期512步结果；点估计不依赖HAC带宽，标准误仍属未校准的探索性结果。',
        table(nino_summary),
        '这不是官方ENSO事件分类、完整大气环流型或外生随机状态；指数为事后发表值，不保证当时实时可获得。外部定义降低直接使用当前H/C响应划分状态的问题，但不能证明大气—海洋系统无共同驱动。',
        '## CERES单月跨产品对照',
        '2019年1月、固定南北大西洋镜像区域各31天、各36个ERA5目标格。CERES先以球面面积重叠映射到相同目标格，再计算相同空间支持的原物理云量面积均值。保留24名义小时槽与四半点槽两种日采样方案，均对ERA5每日四整点样本；原始时间没有改写。',
        table(ceres_table),
        'CERES相对ERA5的正偏差在两套方案均存在，但不能从单月两区域推断全球或长期偏差。卫星与再分析的采样、检索/参数化定义和时间支持不同，相关性不是因果图准确率；这里没有进行季节对照或显著性检验，也没有完成长期独立验证。',
        '## 对审稿意见的实际贡献与剩余限制',
        'R3-6的共同大气驱动控制已增加了具体可复现的实验，并新增同一固定样本上的反方向/反时序诊断；不能据此声称未来信息关联已全部消除或物理方向已识别。R2-3新增不依赖当前H/C定义的Niño3.4外部状态敏感性，但仍非外生干预。CERES完成了小样本数据处理和跨产品描述性对照，未替代长期独立观测验证。',
        '全图有限样本FDR校准、完整水汽预算/轨迹、足够匹配的替代中介、全面地下层修正、跨来源重叠校准仍未完成。700 hPa温度加原850 hPa温度不等于已计算EIS/LTS；六字段也不覆盖辐射、气溶胶及完整垂直结构。新增控制实验限于已冻结样本，不能外推为所有2,574候选滞后的稳健性。正文、最终英文逐条回复和高亮稿仍需将这些证据准确纳入。',
        '## 实施质量与独立复核',
        '物理控制的四张主结果表共8,880行，全部通过行数、唯一键、有效值、样本窗口、输入/代码/输出哈希与汇总算术核验。另以独立最小二乘/SVD和日历HAC复算12组系数与36组分期冻结误差，最大绝对差分别为6.47×10⁻¹³和1.61×10⁻¹⁴。此为全量结构与抽样数值复核，不是每个模型均由第二套程序重跑。五项针对性单元测试通过。',
        '控制设计全部满秩，但方向/路径最高条件数分别约6.20×10⁵和6.27×10⁵，存在强近共线性。最高条件数样本通过QR与SVD独立核对，仍须谨慎解释单项控制系数及其稳定性；满秩不能证明没有共线性，也不代表推断已校准。',
        'CERES的全部62条日配对及四组指标已从原NetCDF独立重算；Niño分析完成432组分状态OLS与12组日历Bartlett计算核对。审计记录、执行快照和全部正反结果一并保留。',
        '## 复核入口',
        '`revision_outputs/physical_controls_inputs/`：预处理计划、训练参数、源清单、直接复算和地面气压提示。\n\n`revision_outputs/physical_controls_experiments/`：全部方向/反时序/路径系数、冻结误差、来源分段、对照变化和执行快照；independent_audit.json记录独立复核。\n\n`revision_outputs/external_nino_regimes/`：完整1,296行交互结果、三个HAC带宽、状态与原生月份统计及独立复算。\n\n`revision_outputs/ceres_pilot_comparison/`：62条日配对、两方案指标、PNG/PDF、保留原始时间的派生文件及独立审计。\n\n`causal-weathergraph/revision/REPRODUCE.md`：运行命令与范围说明。',
        f'生成时间：{datetime.now(timezone.utc).isoformat()}。物理控制模型批次耗时{run["elapsed_seconds"]:.2f}秒；本报告不把运行速度视为算法性能排名。'
    ]
    report = '\n\n'.join(sections) + '\n'
    target = MAJOR / '新增数据处理与补充实验结果_20260930.md'
    target.write_text(report, encoding='utf-8')
    (ROOT / 'revision/NEW_DATA_EXPERIMENTS_ZH.md').write_text(report, encoding='utf-8')
    from docx import Document
    from docx.shared import Pt, Inches, Mm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Mm(210), Mm(297)
    page.top_margin = page.bottom_margin = Inches(0.65)
    page.left_margin = page.right_margin = Inches(0.7)
    style = document.styles['Normal']
    style.font.name, style.font.size = 'Calibri', Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.1
    document.styles['Title'].font.size = Pt(20)
    document.styles['Heading 1'].font.size = Pt(14)
    document.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    document.styles['Heading 1'].paragraph_format.space_after = Pt(5)
    for section in sections:
        if section.startswith('# '):
            document.add_heading(section[2:], 0)
        elif section.startswith('## '):
            document.add_heading(section[3:], 1)
        elif section.startswith('| '):
            rows = section.splitlines()
            cells = [r.strip('| ').split(' | ') for r in rows if not r.startswith('| ---')]
            grid = document.add_table(rows=1, cols=len(cells[0]))
            grid.style = 'Light Shading Accent 1'
            repeat = OxmlElement('w:tblHeader')
            grid.rows[0]._tr.get_or_add_trPr().append(repeat)
            for i, value in enumerate(cells[0]):
                grid.rows[0].cells[i].text = value
            for row in cells[1:]:
                added = grid.add_row().cells
                for cell_, value in zip(added, row):
                    cell_.text = value
            for row in grid.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell_ in row.cells:
                    for paragraph in cell_.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.line_spacing = 1.0
                        for run_ in paragraph.runs:
                            run_.font.size = Pt(9)
        elif section.startswith('生成时间：'):
            footer = page.footer.paragraphs[0]
            footer.text = '2026-09-30 · 新数据补充实验 · '
            number = OxmlElement('w:fldSimple')
            number.set(qn('w:instr'), 'PAGE')
            footer._p.append(number)
        else:
            for paragraph in section.split('\n\n'):
                document.add_paragraph(paragraph.replace('`', ''))
    document.save(target.with_suffix('.docx'))
    print(target)


if __name__ == '__main__':
    main()
